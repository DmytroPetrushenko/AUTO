import hashlib
import math
from datetime import date, datetime
from io import BytesIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from sqlalchemy.exc import IntegrityError
from sqlmodel import select

from auto.constants import normalize
from auto.history import IDENTIFIERS, _digest, duplicate_matches, preview_vehicle
from auto.models import ImportBatch, ImportReceipt, Personnel, Vehicle
from auto.records import require_reader, validate_upload
from auto.schemas import VehicleInput
from auto.services import (
    BusinessError,
    ConflictError,
    list_lookup,
    list_vehicles,
    require_editor,
    save_vehicle,
)
from auto.transactions import Session

REFERENCE_FIELDS = {
    "brand": "vehicle_brands",
    "model": "vehicle_models",
    "vehicle_type": "vehicle_types",
    "support_service": "support_services",
    "unit": "units",
    "engine_type": "engine_types",
    "fuel_type": "fuel_types",
}
INPUT_FIELDS = (
    list(REFERENCE_FIELDS)
    + ["mvo"]
    + [f for f in VehicleInput.model_fields if not f.endswith("_id")]
    + ["vehicle_id"]
)
HEADERS = {
    "brand": "Марка",
    "model": "Модель",
    "vehicle_type": "Тип техніки",
    "support_service": "Служба забезпечення",
    "unit": "Підрозділ",
    "mvo": "МВО",
    "work_unit": "Одиниця напрацювання",
    "military_number": "Військовий номер",
    "vin": "VIN",
    "chassis_number": "Шасі / рама",
    "serial_number": "Заводський номер",
    "accounting_status": "Обліковий статус",
    "technical_condition": "Технічний стан",
    "vehicle_id": "ID автомобіля",
}


def _cell_value(value):
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    return value


def workbook_bytes(rows, columns):
    w = Workbook()
    s = w.active
    s.title = "AUTO"
    s.append(list(columns))
    for row_index, row in enumerate(rows, 2):
        for col, key in enumerate(columns, 1):
            value = _cell_value(row.get(key))
            if isinstance(value, (dict, list)):
                value = str(value)
            cell = s.cell(row_index, col, value)
            # All exported text is literal, including strings starting with '='.
            if isinstance(value, str):
                cell.data_type = "s"
            elif isinstance(value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd"
    for c in s[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="24476A")
    s.freeze_panes = "A2"
    s.auto_filter.ref = s.dimensions
    for col in s.columns:
        s.column_dimensions[col[0].column_letter].width = min(
            45, max(16, len(str(col[0].value or "")) + 2)
        )
    out = BytesIO()
    w.save(out)
    return out.getvalue()


def import_template():
    columns = [HEADERS.get(f, f) for f in INPUT_FIELDS]
    return workbook_bytes([], columns)


def export_full(engine, actor, *, vehicles=None, all_records=False):
    require_reader(actor)
    if all_records:
        vehicles = list_vehicles(engine, include_archived=True)
    if vehicles is None:
        raise BusinessError("Явно оберіть вибірку або всі записи")
    lookups = {
        kind: {r.id: r.name for r in list_lookup(engine, kind, active_only=False)}
        for kind in REFERENCE_FIELDS.values()
    }
    with Session(engine) as session:
        people = {
            p.id: " ".join(filter(None, [p.last_name, p.first_name, p.patronymic]))
            for p in session.exec(select(Personnel)).all()
        }
    rows = []
    for v in vehicles:
        row = v.model_dump()
        for field, kind in REFERENCE_FIELDS.items():
            key = field + "_id"
            row[field] = lookups[kind].get(getattr(v, key))
        row["mvo"] = people.get(v.mvo_id)
        row["engine_type"] = lookups["engine_types"].get(v.engine_type_id)
        row["fuel_type"] = lookups["fuel_types"].get(v.fuel_type_id)
        rows.append(row)
    columns = (
        list(rows[0]) if rows else list(Vehicle.model_fields) + list(REFERENCE_FIELDS) + ["mvo"]
    )
    return workbook_bytes(rows, columns)


def export_screen(frame, actor):
    require_reader(actor)
    return workbook_bytes(frame.to_dict("records"), list(frame.columns))


def _load(content, *, data_only):
    try:
        return load_workbook(BytesIO(content), data_only=data_only)
    except (
        ValueError,
        KeyError,
        TypeError,
        IndexError,
        BadZipFile,
        ParseError,
        InvalidFileException,
    ) as exc:
        raise BusinessError("Пошкоджений або непідтримуваний XLSX") from exc


def inspect_workbook(content):
    validate_upload("import.xlsx", content)
    raw = _load(content, data_only=False)
    cache = _load(content, data_only=True)
    result = []
    for sheet in raw:
        formulas = [
            (c.coordinate, c.value, cache[sheet.title][c.coordinate].value)
            for row in sheet
            for c in row
            if c.data_type == "f"
        ]
        result.append(
            {
                "sheet": sheet.title,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "formulas": len(formulas),
                "missing_cache": sum(v is None for _, _, v in formulas),
            }
        )
    raw.close()
    cache.close()
    return result


def headers(content, sheet, header_row=1):
    validate_upload("import.xlsx", content)
    w = _load(content, data_only=False)
    try:
        if not 1 <= header_row <= 50 or sheet not in w.sheetnames:
            raise BusinessError("Неправильний аркуш або рядок заголовків")
        s = w[sheet]
        return {
            str(c.column): str(c.value) if c.value is not None else f"Стовпець {c.column}"
            for c in s[header_row]
            if c.column <= 200
        }
    finally:
        w.close()


class ImportRowError(BusinessError):
    def __init__(self, issues):
        self.issues = list(dict.fromkeys(issues))
        super().__init__("; ".join(self.issues))


def _resolve(engine, values):
    issues = []
    data = {k: v for k, v in values.items() if k in VehicleInput.model_fields}
    for field in list(data):
        definition = VehicleInput.model_fields[field]
        if data[field] is None and not definition.is_required() and definition.default is not None:
            data.pop(field)
    for field, kind in REFERENCE_FIELDS.items():
        value = values.get(field)
        if value is None or str(value).strip() == "":
            continue
        matches = [
            r for r in list_lookup(engine, kind) if normalize(r.name) == normalize(str(value))
        ]
        if field == "model":
            matches = [r for r in matches if r.brand_id == data.get("brand_id")]
        if len(matches) != 1:
            issues.append(
                f"Довідник {HEADERS.get(field, field)}: "
                f"немає однозначного активного значення «{value}»"
            )
        else:
            data[field + "_id"] = matches[0].id
    if values.get("mvo"):
        if not isinstance(values["mvo"], str):
            issues.append("МВО має бути текстовим ПІБ")
        elif isinstance(values["mvo"], str):
            with Session(engine) as session:
                matches = [
                    p
                    for p in session.exec(select(Personnel)).all()
                    if p.is_active
                    and normalize(" ".join(filter(None, [p.last_name, p.first_name, p.patronymic])))
                    == normalize(values["mvo"])
                ]
            if len(matches) != 1:
                issues.append("МВО: ПІБ не має однозначного активного збігу")
            else:
                data["mvo_id"] = matches[0].id
    for f in IDENTIFIERS:
        value = data.get(f)
        if value is None:
            continue
        if not isinstance(value, str):
            issues.append(f"{f}: числовий ідентифікатор потребує ручного перегляду")
        elif normalize(value) in ("б/н", "бн", "----", "-", "невідомо"):
            issues.append(f"{f}: неоднозначне позначення замість ідентифікатора")
    if (
        isinstance(data.get("accounting_status"), str)
        and normalize(data["accounting_status"]) == "вибула"
    ):
        data["accounting_status"] = "Передано"
    try:
        result = VehicleInput.model_validate(data).model_dump()
    except ValidationError as exc:
        for error in exc.errors(include_url=False, include_input=False):
            field = str(error["loc"][0]) if error["loc"] else "Картка"
            if field in IDENTIFIERS and any(i.startswith(field + ":") for i in issues):
                continue
            issues.append(f"{field}: {error['msg']}")
    if issues:
        raise ImportRowError(issues)
    return result


def preview_import(
    engine,
    content,
    mapping,
    actor,
    *,
    sheet="AUTO",
    header_row=1,
    defaults=None,
    use_cached=False,
    overrides=None,
    source_name="import.xlsx",
    required_explicit=None,
):
    require_editor(actor)
    validate_upload("import.xlsx", content)
    if set(mapping) - set(INPUT_FIELDS):
        raise BusinessError("Невідомі поля мапінгу")
    if set(required_explicit or []) - set(INPUT_FIELDS):
        raise BusinessError("Невідомі поля явного погодження")
    raw = _load(content, data_only=False)
    cache = _load(content, data_only=True)
    if sheet not in raw.sheetnames or not 1 <= header_row <= 50:
        raise BusinessError("Неправильний аркуш/заголовок")
    s = raw[sheet]
    if s.max_row > 10000 or s.max_column > 200:
        raise BusinessError("Ліміт імпорту: 10000 рядків і 200 стовпців")
    colmap = {k: int(v) for k, v in mapping.items() if v}
    if any(not 1 <= c <= 200 for c in colmap.values()):
        raise BusinessError("Неправильний стовпець")
    rows = []
    seen = {}
    seen_fingerprints = set()
    try:
        for index in range(header_row + 1, s.max_row + 1):
            if not any(s.cell(index, c).value is not None for c in colmap.values()):
                continue
            values = dict(defaults or {})
            errors = []
            formulas = []
            source_cells = {}
            unchecked = {}
            for field, col in colmap.items():
                cell = s.cell(index, col)
                source_cells[field] = {
                    "cell": cell.coordinate,
                    "type": cell.data_type,
                    "format": cell.number_format,
                }
                value = cell.value
                if cell.data_type == "f":
                    cached = cache[sheet][cell.coordinate]
                    formulas.append(
                        {"cell": cell.coordinate, "formula": value, "cached": str(cached.value)}
                    )
                    if not use_cached:
                        errors.append(f"{cell.coordinate}: формула потребує явного вибору кешу")
                        unchecked[field] = "Формула без погодженого використання кешу"
                    if cached.value is None or cached.data_type == "e":
                        errors.append(f"{cell.coordinate}: відсутній/помилковий кеш")
                        unchecked[field] = "Відсутній/помилковий кеш"
                    value = cached.value
                if cell.data_type == "e":
                    errors.append(f"{cell.coordinate}: помилка Excel")
                values[field] = _cell_value(value)
            values.update((overrides or {}).get(str(index), {}))
            # Coordinates are never an import target; extra source columns remain untouched.
            values = {k: v for k, v in values.items() if k in INPUT_FIELDS}
            for field in required_explicit or []:
                value = values.get(field)
                if value is None or (isinstance(value, str) and not value.strip()):
                    errors.append(f"{field}: потрібне явне значення; типове не є фактом джерела")
            data = None
            plan = None
            incoming = []
            duplicates = []
            checked = []
            # Diagnostics must not disappear because another column is invalid.
            # Numeric/formula/placeholder identifiers remain unconfirmed, never auto-coerced.
            with Session(engine) as session:
                for f in IDENTIFIERS:
                    value = values.get(f)
                    if f in unchecked:
                        continue
                    if value is None or value == "":
                        continue
                    if not isinstance(value, str):
                        unchecked[f] = "Числове подання потребує ручного погодження"
                        continue
                    n = normalize(value)
                    if not n or n in ("б/н", "бн", "----", "-", "невідомо"):
                        unchecked[f] = "Порожній/неоднозначний ідентифікатор"
                        continue
                    checked.append(f)
                    matches = duplicate_matches(session, f, value)
                    if matches:
                        duplicates.append(
                            {
                                "field": f,
                                "value": value,
                                "matches": matches,
                                "approval_required": True,
                            }
                        )
                    key = (f, n)
                    if key in seen:
                        incoming.append({"field": f, "source_row": seen[key]})
                    else:
                        seen[key] = index
            fingerprint = None
            try:
                data = _resolve(engine, values)
                fingerprint = _digest(data)
                with Session(engine) as session:
                    if values.get("vehicle_id") and session.get(Vehicle, str(values["vehicle_id"])):
                        errors.append("ID уже існує; оновлення карток через імпорт заборонене")
                    if session.exec(
                        select(ImportReceipt).where(ImportReceipt.fingerprint == fingerprint)
                    ).first():
                        errors.append("Цей рядок уже імпортований; наявну картку не змінено")
                if fingerprint in seen_fingerprints:
                    errors.append("Ідентичний рядок уже є у вибірці; потрібне ручне рішення")
                seen_fingerprints.add(fingerprint)
                plan = preview_vehicle(engine, data, actor)
            except ImportRowError as exc:
                errors.extend(exc.issues)
            except (BusinessError, ValidationError, PermissionError) as exc:
                errors.append(str(exc))
            rows.append(
                {
                    "row": index,
                    "source": {"file": source_name, "sheet": sheet, "row": index},
                    "source_cells": source_cells,
                    "input": values,
                    "data": data,
                    "errors": errors,
                    "formulas": formulas,
                    "duplicates": duplicates,
                    "duplicate_checked_fields": checked,
                    "duplicate_unchecked_fields": {
                        f: reason for f, reason in unchecked.items() if f in IDENTIFIERS
                    },
                    "incoming": incoming,
                    "fingerprint": fingerprint,
                    "plan_token": plan["token"] if plan else None,
                    "applied_defaults": {
                        f: data[f]
                        for f, definition in VehicleInput.model_fields.items()
                        if data is not None
                        and not definition.is_required()
                        and definition.default is not None
                        and values.get(f) is None
                    },
                }
            )
    finally:
        raw.close()
        cache.close()
    settings = {
        "mapping": mapping,
        "sheet": sheet,
        "header_row": header_row,
        "defaults": defaults or {},
        "use_cached": use_cached,
        "overrides": overrides or {},
        "source_name": source_name,
        "required_explicit": list(required_explicit or []),
    }
    preview = {
        "rows": rows,
        "settings": settings,
        "source_hash": hashlib.sha256(content).hexdigest(),
        "actor": {"id": actor.user_id, "role": actor.role.value},
    }
    preview["token"] = _digest(preview)
    preview["_content"] = content
    return preview


def apply_import(engine, preview, actor, *, confirmed=False, approvals=None):
    require_editor(actor)
    if not confirmed:
        raise BusinessError("Потрібне явне підтвердження перегляду імпорту")
    settings = preview["settings"]
    refreshed = preview_import(
        engine,
        preview["_content"],
        settings["mapping"],
        actor,
        **{k: v for k, v in settings.items() if k != "mapping"},
    )
    if refreshed["token"] != preview["token"]:
        raise ConflictError("Перегляд імпорту застарів; перевірте його повторно")
    with Session(engine) as session:
        batch = ImportBatch(
            source_name=settings["source_name"],
            source_hash=preview["source_hash"],
            user_id=actor.user_id,
        )
        session.add(batch)
        session.commit()
        session.refresh(batch)
    results = []
    created_ids = set()
    for row in refreshed["rows"]:
        if row["errors"]:
            results.append(
                {"row": row["row"], "status": "skipped", "reason": " | ".join(row["errors"])}
            )
            continue
        try:
            p = preview_vehicle(engine, row["data"], actor)
            allowed = {
                m["vehicle_id"] for d in row["duplicates"] for m in d["matches"]
            } | created_ids
            actual = {m["vehicle_id"] for d in p["duplicates"] for m in d["matches"]}
            if not actual <= allowed:
                raise ConflictError("З’явилися нові дублікати; повторіть перегляд")
            vehicle = save_vehicle(
                engine,
                row["data"],
                actor,
                preview_token=p["token"],
                confirmed=True,
                approvals=(approvals or {}).get(str(row["row"]), {}),
                import_receipt={
                    "batch_id": batch.id,
                    "source_row": row["row"],
                    "fingerprint": row["fingerprint"],
                },
            )
            created_ids.add(vehicle.id)
            results.append({"row": row["row"], "status": "created", "vehicle_id": vehicle.id})
        except (BusinessError, ValidationError, PermissionError, IntegrityError) as exc:
            results.append({"row": row["row"], "status": "skipped", "reason": str(exc)})
    with Session(engine) as session:
        saved = session.get(ImportBatch, batch.id)
        saved.results = results
        session.add(saved)
        session.commit()
    return results


def import_log(engine, actor):
    require_reader(actor)
    with Session(engine) as s:
        return s.exec(select(ImportBatch).order_by(ImportBatch.created_at.desc())).all()
