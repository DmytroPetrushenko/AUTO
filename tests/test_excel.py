from datetime import date
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from auto.excel import (
    HEADERS,
    apply_import,
    export_full,
    export_screen,
    headers,
    import_log,
    import_template,
    preview_import,
    workbook_bytes,
)
from auto.services import (
    Actor,
    BusinessError,
    ConflictError,
    Role,
    get_vehicle,
    list_lookup,
    list_vehicles,
    save_vehicle,
)

VIEWER = Actor("reader", Role.VIEWER)


def source(engine, minimum):
    return {
        "brand": list_lookup(engine, "vehicle_brands")[0].name,
        "model": list_lookup(engine, "vehicle_models")[0].name,
        "vehicle_type": next(
            t.name
            for t in list_lookup(engine, "vehicle_types")
            if t.id == minimum["vehicle_type_id"]
        ),
        "support_service": next(
            t.name
            for t in list_lookup(engine, "support_services")
            if t.id == minimum["support_service_id"]
        ),
        "work_unit": "км",
    }


def book(rows):
    columns = list(dict.fromkeys(k for row in rows for k in row))
    return workbook_bytes(rows, columns), {f: str(i + 1) for i, f in enumerate(columns)}


def test_screen_export_order_types_blanks_and_formula_safety():
    frame = pd.DataFrame(
        [
            {"Name": '=HYPERLINK("bad")', "Mass": 10.5, "Date": date(2025, 1, 1), "Missing": None},
            {"Name": "B", "Mass": 0, "Date": None, "Missing": None},
        ]
    )
    data = export_screen(frame[["Date", "Mass", "Name", "Missing"]], VIEWER)
    s = load_workbook(BytesIO(data)).active
    assert [c.value for c in s[1]] == ["Date", "Mass", "Name", "Missing"]
    assert s["A2"].is_date and s["B2"].data_type == "n" and s["B3"].value == 0
    assert s["D2"].value is None
    assert s["C2"].data_type == "s" and s["C2"].value.startswith("=")


def test_full_export_selection_all_and_template(engine, minimum, actor):
    v = save_vehicle(engine, minimum | {"seats": 3, "accounting_date": date(2025, 1, 1)}, actor)
    save_vehicle(engine, minimum | {"accounting_status": "Списано"}, actor)
    s = load_workbook(BytesIO(export_full(engine, VIEWER, vehicles=[v]))).active
    names = [c.value for c in s[1]]
    assert set(type(v).model_fields) <= set(names)
    assert s.cell(2, names.index("seats") + 1).value == 3
    assert s.cell(2, names.index("accounting_date") + 1).is_date
    assert load_workbook(BytesIO(export_full(engine, VIEWER, all_records=True))).active.max_row == 3
    with pytest.raises(BusinessError):
        export_full(engine, VIEWER)
    assert HEADERS["brand"] in headers(import_template(), "AUTO").values()


def test_preview_no_write_import_skips_errors_logs_and_repeat(engine, minimum, actor):
    row = source(engine, minimum)
    data, mapping = book(
        [
            row | {"military_number": "ONE"},
            row | {"military_number": "BAD", "payload_kg": -1},
            row | {"military_number": "THREE", "accounting_status": "ВИБУЛА"},
        ]
    )
    p = preview_import(engine, data, mapping, actor)
    assert len(p["rows"]) == 3 and p["rows"][1]["errors"]
    assert list_vehicles(engine, include_archived=True) == []
    with pytest.raises(BusinessError):
        apply_import(engine, p, actor)
    result = apply_import(engine, p, actor, confirmed=True)
    assert [r["status"] for r in result] == ["created", "skipped", "created"]
    assert list_vehicles(engine, accounting_status="Передано")[0].unit_id is None
    before = [(v.id, v.version) for v in list_vehicles(engine, include_archived=True)]
    again = apply_import(
        engine, preview_import(engine, data, mapping, actor), actor, confirmed=True
    )
    assert all(r["status"] == "skipped" for r in again)
    assert [(v.id, v.version) for v in list_vehicles(engine, include_archived=True)] == before
    assert len(import_log(engine, VIEWER)) == 2


def test_arbitrary_mapping_manual_numeric_id_and_coordinates_ignored(engine, minimum, actor):
    row = source(engine, minimum)
    data, _ = book(
        [
            {
                "b": row["brand"],
                "m": row["model"],
                "t": row["vehicle_type"],
                "s": row["support_service"],
                "u": "км",
                "N": 123,
                "Coordinates": "DO NOT IMPORT",
            }
        ]
    )
    mapping = dict(
        zip(
            ["brand", "model", "vehicle_type", "support_service", "work_unit", "military_number"],
            map(str, range(1, 7)),
            strict=True,
        )
    )
    p = preview_import(engine, data, mapping, actor)
    assert "числовий" in p["rows"][0]["errors"][0]
    p = preview_import(engine, data, mapping, actor, overrides={"2": {"military_number": "00123"}})
    assert not p["rows"][0]["errors"]
    result = apply_import(engine, p, actor, confirmed=True)
    vehicle = get_vehicle(engine, result[0]["vehicle_id"])
    assert vehicle.military_number == "00123" and vehicle.current_location is None


def test_formulas_separate_and_cache_missing_blocks(engine, minimum, actor):
    row = source(engine, minimum)
    data, mapping = book([row | {"military_number": "A", "payload_kg": 100}])
    w = load_workbook(BytesIO(data))
    w.active.cell(2, int(mapping["payload_kg"]), "=50+50")
    stream = BytesIO()
    w.save(stream)
    p = preview_import(engine, stream.getvalue(), mapping, actor)
    assert p["rows"][0]["formulas"] and p["rows"][0]["errors"]
    p = preview_import(engine, stream.getvalue(), mapping, actor, use_cached=True)
    assert any("кеш" in e for e in p["rows"][0]["errors"])


def test_import_duplicate_approvals_and_stale_preview(engine, minimum, actor):
    row = source(engine, minimum)
    save_vehicle(engine, minimum | {"military_number": "SAME"}, actor)
    data, mapping = book([row | {"military_number": "same"}])
    p = preview_import(engine, data, mapping, actor)
    assert p["rows"][0]["duplicates"]
    result = apply_import(engine, p, actor, confirmed=True)
    assert result[0]["status"] == "skipped"
    p = preview_import(engine, data, mapping, actor)
    result = apply_import(
        engine,
        p,
        actor,
        confirmed=True,
        approvals={"2": {"military_number": {"confirmed": True, "reason": "Allowed source"}}},
    )
    assert result[0]["status"] == "created"
    p = preview_import(engine, *book([row | {"military_number": "FRESH"}]), actor)
    save_vehicle(engine, minimum | {"military_number": "FRESH"}, actor)
    with pytest.raises(ConflictError):
        apply_import(engine, p, actor, confirmed=True)
    with pytest.raises(PermissionError):
        preview_import(engine, data, mapping, VIEWER)


def test_incoming_duplicate_has_same_approval_rules(engine, minimum, actor):
    row = source(engine, minimum)
    data, mapping = book(
        [
            row | {"military_number": "SAME", "color": "A"},
            row | {"military_number": "SAME", "color": "B"},
        ]
    )
    p = preview_import(engine, data, mapping, actor)
    assert p["rows"][1]["incoming"] == [{"field": "military_number", "source_row": 2}]
    result = apply_import(
        engine,
        p,
        actor,
        confirmed=True,
        approvals={"3": {"military_number": {"confirmed": True, "reason": "Separate cards"}}},
    )
    assert [r["status"] for r in result] == ["created", "created"]


def test_identifier_free_import_is_idempotent(engine, minimum, actor):
    row = source(engine, minimum)
    data, mapping = book([row])
    p = preview_import(engine, data, mapping, actor)
    apply_import(engine, p, actor, confirmed=True)
    assert len(list_vehicles(engine)) == 1
    # An identifier-free row is still idempotent; it never overwrites or silently duplicates.
    again = preview_import(engine, data, mapping, actor)
    assert any("уже імпортований" in error for error in again["rows"][0]["errors"])


def test_import_receipt_failure_rolls_back_vehicle(engine, minimum, actor):
    from sqlalchemy.exc import IntegrityError

    with pytest.raises(IntegrityError):
        save_vehicle(
            engine,
            minimum,
            actor,
            import_receipt={"batch_id": "missing-batch", "source_row": 2, "fingerprint": "unique"},
        )
    assert list_vehicles(engine) == []
